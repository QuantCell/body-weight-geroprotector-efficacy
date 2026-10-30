"""Summarize comparator survival and population context."""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
import csv
import json
import xml.etree.ElementTree as ET

from lifelines import KaplanMeierFitter
from openpyxl import load_workbook
from scipy import stats
import numpy as np
import pandas as pd

from data_utils import clean_adult_weights


def summarize_comparator_survival():
    ROOT = Path(__file__).absolute().parents[2]

    LIFE = ROOT / "data/inputs/ITP_bodyweight_all_cohorts.xlsx"

    CELLS = ROOT / "data/derived/comparator_effect_cells.csv"

    SUMMARY_OUT = ROOT / "data/derived/comparator_state.csv"

    CELL_OUT = ROOT / "data/derived/comparator_cell_medians.csv"

    Z95 = 1.959963984540054

    LANDMARK_DAYS = 6 * 30.4

    def km_median(frame):
        """Six-month-left-truncated Kaplan-Meier median."""
        z = frame.loc[frame["lifespan_days"].gt(LANDMARK_DAYS),
                      ["lifespan_days", "dead"]].dropna()
        if z.empty:
            return np.nan
        fit = KaplanMeierFitter()
        fit.fit(
            z["lifespan_days"],
            event_observed=z["dead"],
            entry=np.full(len(z), LANDMARK_DAYS),
        )
        return float(fit.median_survival_time_)

    life = pd.read_excel(LIFE)

    for column in ("lifespan_days", "dead"):
        life[column] = pd.to_numeric(life[column], errors="coerce")

    life = life[life.lifespan_days.notna() & life.dead.notna()].copy()

    cells = pd.read_csv(CELLS)

    if len(cells) != 534:
        raise AssertionError(f"Expected 534 comparator cells, found {len(cells)}")

    medians = []

    for row in cells.itertuples(index=False):
        treated = life[(life.cohort == row.cohort) & (life.group == row.group) &
                       (life.site == row.site) & (life.sex == row.sex)]
        control = life[(life.cohort == row.cohort) & (life.group == "Control") &
                       (life.site == row.site) & (life.sex == row.sex)]
        cmed, tmed = km_median(control), km_median(treated)
        medians.append((cmed, tmed, 100 * (tmed - cmed) / cmed))

    cells[["control_km_median", "treated_km_median", "median_gain_pct"]] = pd.DataFrame(
        medians, index=cells.index
    )

    cells["control_cluster"] = cells.cohort.astype(str) + "|" + cells.site.astype(str)

    cells["sex_specific_control"] = cells.control_cluster + "|" + cells.sex.astype(str)

    cells.to_csv(CELL_OUT, index=False)

    rows = []

    for sex in ("m", "f"):
        subset = cells[cells.sex == sex].copy()
        for split in (3, 4):
            # Form quantiles of observed control weights separately within sex.
            subset["bin"] = pd.qcut(
                subset.ctl_bw6, split, labels=[f"Q{i}" for i in range(1, split + 1)]
            )
            for label, group in subset.groupby("bin", observed=True):
                weights = 1 / group.se.pow(2)
                beta = np.average(group.log_hr, weights=weights)
                se = np.sqrt(1 / weights.sum())
                rows.append({
                    "sex": sex,
                    "split": split,
                    "bin": str(label),
                    "cells": len(group),
                    "mean_control_weight_g": group.ctl_bw6.mean(),
                    "pooled_hr": np.exp(beta),
                    "lower_95_ci": np.exp(beta - Z95 * se),
                    "upper_95_ci": np.exp(beta + Z95 * se),
                    "hazard_reduction_pct": 100 * (1 - np.exp(beta)),
                    "mean_cell_km_median_gain_pct": group.median_gain_pct.mean(),
                })

    summary = pd.DataFrame(rows)

    summary.to_csv(SUMMARY_OUT, index=False)

    print(f"{len(cells)} cells; {cells.control_cluster.nunique()} cohort-site control clusters; "
          f"{cells.sex_specific_control.nunique()} sex-specific control means")

    print(summary.to_string(index=False))


def summarize_weight_lifespan_correlations():
    PACKAGE = Path(__file__).absolute().parents[2]

    INPUT = PACKAGE / "supplementary" / "Supplementary_Data_1-10.xlsx"

    OUTPUT = PACKAGE / "data" / "derived" / "weight_lifespan_by_experiment.csv"

    data = clean_adult_weights(pd.read_excel(INPUT, sheet_name="SD2_mouse_level"))
    data = data[
        data["dead"].eq(1)
        & data["bw_6"].notna()
        & data["lifespan_days"].notna()
    ].copy()

    rows = []
    for (cohort, sex, group), stratum in data.groupby(
        ["cohort", "sex", "group"], sort=True
    ):
        site_estimates = []
        for _, cell in stratum.groupby("site", sort=True):
            if len(cell) < 12:
                continue
            rho = cell[["bw_6", "lifespan_days"]].corr(method="spearman").iloc[0, 1]
            if not np.isfinite(rho):
                continue
            site_estimates.append(
                (np.arctanh(np.clip(rho, -0.999999, 0.999999)), len(cell) - 3, len(cell))
            )
        if not site_estimates:
            continue
        z = np.asarray([value[0] for value in site_estimates])
        weights = np.asarray([value[1] for value in site_estimates], dtype=float)
        rows.append(
            {
                "cohort": cohort,
                "sex": sex,
                "group": group,
                "nsite": len(site_estimates),
                "n": int(sum(value[2] for value in site_estimates)),
                "rho": float(np.tanh(np.average(z, weights=weights))),
            }
        )

    result = pd.DataFrame(rows)
    result.to_csv(OUTPUT, index=False)

    expected = {"m": (109, 108), "f": (105, 87)}
    for sex, (n_expected, negative_expected) in expected.items():
        subset = result[result["sex"].eq(sex)]
        observed = (len(subset), int(subset["rho"].lt(0).sum()))
        if observed != (n_expected, negative_expected):
            raise RuntimeError(
                f"Reconciliation failed for sex={sex}: {observed} != "
                f"{(n_expected, negative_expected)}"
            )
        print(
            f"{sex}: {observed[1]}/{observed[0]} negative; "
            f"median rho={subset['rho'].median():.3f}"
        )
    print(OUTPUT)


def summarize_sex_survival():
    ROOT = Path(__file__).absolute().parents[2]

    INPUTS = ROOT / "data" / "inputs"

    OUTPUTS = ROOT / "data" / "derived"

    CHENG_COHORTS = {"C2004", "C2005", "C2006", "C2007", "C2009", "C2010"}

    def km_summary(
        frame: pd.DataFrame,
        day: float = 500.0,
        all_events: bool = False,
    ) -> tuple[float, float]:
        """Return the Kaplan-Meier median and cumulative mortality at ``day``."""
        event = pd.Series(1, index=frame.index) if all_events else frame["dead"].astype(int)
        fit = KaplanMeierFitter().fit(frame["age"], event_observed=event)
        median = float(fit.median_survival_time_)
        mortality = 100.0 * (1.0 - float(fit.predict(day)))
        return median, mortality

    def summarize_set(frame: pd.DataFrame, label: str) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        for sex, name in (("m", "male"), ("f", "female")):
            subset = frame[frame["sex"].eq(sex)]
            median, mortality = km_summary(subset)
            _, mortality_all_events = km_summary(subset, all_events=True)
            rows.append({
                "cohort_set": label,
                "sex": name,
                "n": len(subset),
                "cohorts": subset["cohort"].nunique(),
                "km_median_days": median,
                "day500_mortality_pct": mortality,
                "day500_mortality_all_events_pct": mortality_all_events,
            })
        female, male = rows[1], rows[0]
        rows.append({
            "cohort_set": label,
            "sex": "female_minus_male",
            "n": np.nan,
            "cohorts": np.nan,
            "km_median_days": female["km_median_days"] - male["km_median_days"],
            "day500_mortality_pct": (
                female["day500_mortality_pct"] - male["day500_mortality_pct"]
            ),
            "day500_mortality_all_events_pct": np.nan,
        })
        return rows

    controls = pd.read_csv(INPUTS / "itp_gn_earlylife_controls.csv")
    controls = controls[controls["age"].notna()].copy()
    external = pd.read_csv(INPUTS / "sex_specific_external_survival_context.csv")
    OUTPUTS.mkdir(parents=True, exist_ok=True)

    summary = pd.DataFrame(
        summarize_set(controls, "all_17_cohorts")
        + summarize_set(
            controls[controls["cohort"].isin(CHENG_COHORTS)],
            "published_six_cohorts",
        )
    )
    drido = pd.read_csv(OUTPUTS / "drido_180_landmark_pooled_cr.csv")
    al = drido.loc[drido["group"].eq("AL")].rename(columns={"SurvDays": "age"})
    median, mortality = km_summary(al, all_events=True)
    summary = pd.concat([summary, pd.DataFrame([{
        "cohort_set": "DRiDO_ad_libitum", "sex": "female", "n": len(al),
        "cohorts": 1, "km_median_days": median,
        "day500_mortality_pct": mortality,
        "day500_mortality_all_events_pct": mortality,
    }])], ignore_index=True)
    summary.to_csv(OUTPUTS / "sex_survival_context.csv", index=False, lineterminator="\n")

    cohort_rows: list[dict[str, object]] = []
    for cohort, block in controls.groupby("cohort"):
        male = block[block["sex"].eq("m")]
        female = block[block["sex"].eq("f")]
        if male.empty or female.empty:
            continue
        male_median, _ = km_summary(male)
        female_median, _ = km_summary(female)
        cohort_rows.append({
            "cohort": cohort,
            "n": len(block),
            "male_km_median_days": male_median,
            "female_km_median_days": female_median,
            "female_advantage_days": female_median - male_median,
        })
    pd.DataFrame(cohort_rows).sort_values("cohort").to_csv(
        OUTPUTS / "sex_survival_by_cohort.csv", index=False, lineterminator="\n"
    )

    external["female_minus_male_days"] = (
        external["female_survival_days"] - external["male_survival_days"]
    )
    external.to_csv(
        OUTPUTS / "external_sex_survival_context.csv",
        index=False,
        lineterminator="\n",
    )


def summarize_founder_context():
    ROOT = Path(__file__).absolute().parents[2]

    INPUTS = ROOT / "data" / "inputs"

    OUTPUTS = ROOT / "data" / "derived"

    def read_rows(path):
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return list(csv.DictReader(handle))

    founders = read_rows(INPUTS / "founder_strain_mpd_reed1.csv")

    founder_mean = sum(float(row["reported_weight_g"]) for row in founders) / len(founders)

    itp = clean_adult_weights(pd.read_excel(INPUTS / "ITP_bodyweight_all_cohorts.xlsx"))

    itp_male_controls = itp.loc[
        itp["group"].eq("Control") & itp["sex"].eq("m")
        & itp["bw_6"].notna() & itp["dead"].notna()
        & itp["lifespan_days"].gt(6 * 30.4), "bw_6"
    ]

    itp_mean = float(itp_male_controls.mean())

    sensitivities = read_rows(OUTPUTS / "matched_quartile_standardized_survival.csv")

    lowest = [
        row for row in sensitivities
        if row["quartile"] == "1"
    ]

    if len(lowest) != 1:
        raise RuntimeError(f"Expected one lowest-quarter row; found {len(lowest)}")

    lowest_mean = float(lowest[0]["mean_control_weight"])

    output = [
        {"quantity": "four_founder_mean_weight_g", "value": founder_mean},
        {"quantity": "itp_male_control_6m_n", "value": len(itp_male_controls)},
        {"quantity": "itp_male_control_6m_mean_weight_g", "value": itp_mean},
        {"quantity": "itp_minus_founder_mean_g", "value": itp_mean - founder_mean},
        {"quantity": "itp_percent_above_founder_mean", "value": 100 * (itp_mean / founder_mean - 1)},
        {"quantity": "lowest_itp_quarter_mean_weight_g", "value": lowest_mean},
        {"quantity": "lowest_itp_quarter_percent_above_founder_mean",
         "value": 100 * (lowest_mean / founder_mean - 1)},
        {"quantity": "heaviest_founder_reported_mean_weight_g", "value": max(float(row["reported_weight_g"]) for row in founders)},
    ]

    out_path = OUTPUTS / "founder_strain_weight_comparison.csv"

    with out_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["quantity", "value"])
        writer.writeheader()
        writer.writerows(output)

    print(out_path)

    def strain_means(path, sex):
        raw = json.loads(path.read_text(encoding="utf-8"))
        rows = raw.get("strainmeans", raw) if isinstance(raw, dict) else raw
        return {
            row["strain"]: float(row["mean"])
            for row in rows
            if str(row.get("sex", "")).lower().startswith(sex)
            and row.get("mean") is not None
        }

    weights = strain_means(INPUTS / "mpd_32603_strainmeans.json", "m")

    ages = strain_means(INPUTS / "mpd_32680_strainmeans.json", "m")

    age_matched = {strain: value for strain, value in weights.items()
                   if strain in ages and ages[strain] <= 36.0}

    male_fat = strain_means(INPUTS / "mpd_10331_strainmeans.json", "m")

    female_fat = strain_means(INPUTS / "mpd_10331_strainmeans.json", "f")

    workbook = load_workbook(ROOT / "supplementary" / "Supplementary_Data_1-10.xlsx",
                             read_only=True, data_only=True)

    sheet = workbook["SD5_body_composition"]

    rows = sheet.iter_rows(values_only=True)

    headers = next(rows)

    body = [dict(zip(headers, row)) for row in rows]

    workbook.close()

    itp_male_fat = next(float(row["bodyfat_pct_control"]) for row in body
                        if row["strain"] == "UM-HET3" and row["sex"] == "Male"
                        and row["Intervention"] == "Rapamycin" and "42 ppm" in str(row["dosage"]))

    itp_female_fat = next(float(row["bodyfat_pct_control"]) for row in body
                          if row["strain"] == "UM-HET3" and row["sex"] == "Female"
                          and row["Intervention"] == "Rapamycin" and "42 ppm" in str(row["dosage"]))

    strain_context = [
        {"quantity": "age_matched_strains", "value": len(age_matched)},
        {"quantity": "survey_mean_weight_g",
         "value": sum(age_matched.values()) / len(age_matched)},
        {"quantity": "weight_survey_strains", "value": len(age_matched)},
        {"quantity": "weight_strains_at_or_above_itp_mean",
         "value": sum(value >= itp_mean for value in age_matched.values())},
        {"quantity": "weight_strains_at_or_above_lightest_quarter",
         "value": sum(value >= lowest_mean for value in age_matched.values())},
        {"quantity": "male_fat_strains", "value": len(male_fat)},
        {"quantity": "male_fat_at_or_above_itp",
         "value": sum(value >= itp_male_fat for value in male_fat.values())},
        {"quantity": "female_fat_strains", "value": len(female_fat)},
        {"quantity": "female_fat_at_or_above_itp",
         "value": sum(value >= itp_female_fat for value in female_fat.values())},
    ]

    strain_path = OUTPUTS / "strain_context_summary.csv"

    with strain_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["quantity", "value"])
        writer.writeheader()
        writer.writerows(strain_context)

    print(strain_path)


def summarize_control_lifespan():
    ROOT = Path(__file__).absolute().parents[2]

    SOURCE = ROOT / "supplementary" / "Supplementary_Data_1-10.xlsx"

    OUTPUT = ROOT / "data" / "derived" / "control_lifespan_ceiling.csv"

    experiments = pd.read_excel(SOURCE, sheet_name="SD10_control_lifespan")

    induced = experiments.lifespan_evidence_quote.fillna("").str.contains("doxorubicin", case=False)

    experiments = experiments[~induced]

    experiments = experiments[experiments.sex.isin(["Female", "Male"])].copy()

    experiments["is_itp"] = experiments.itp_status.astype(str).str.upper().eq("ITP")

    def control_groups(frame):
        return frame.drop_duplicates(["pmid", "lifespan_control_days"])

    def study_medians(frame):
        return frame.groupby("pmid", as_index=False).lifespan_control_days.median()

    specifications = [
        ("unique_control_groups", control_groups),
        ("study_level_median", study_medians),
        ("median_endpoints_only",
         lambda frame: control_groups(frame[frame.lifespan_measure.eq("median")])),
        ("published_2000_or_later", lambda frame: control_groups(frame[frame.year.ge(2000)])),
    ]

    rows = []

    for label, select in specifications:
        for sex in ("Female", "Male"):
            itp = select(experiments[experiments.is_itp & experiments.sex.eq(sex)]).lifespan_control_days
            non_itp = select(experiments[~experiments.is_itp & experiments.sex.eq(sex)]).lifespan_control_days
            itp_median = itp.median()
            rows.append({
                "specification": label,
                "sex": sex,
                "itp_n": len(itp),
                "itp_median_days": itp_median,
                "non_itp_n": len(non_itp),
                "non_itp_median_days": non_itp.median(),
                "non_itp_reaching_itp_median": int((non_itp >= itp_median).sum()),
                "p_mannwhitney": stats.mannwhitneyu(itp, non_itp, alternative="two-sided").pvalue,
            })

    pd.DataFrame(rows).to_csv(OUTPUT, index=False, lineterminator="\n")


def verify_strain_survival_summary():
    root = Path(__file__).absolute().parents[2]
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(root / "data/inputs/yuan2020_supplementary_tables.docx") as archive:
        document = ET.fromstring(archive.read("word/document.xml"))
    rows = []
    for row in document.findall(".//w:tbl", ns)[0].findall("w:tr", ns)[2:]:
        cells = ["".join(node.text or "" for node in cell.findall(".//w:t", ns))
                 for cell in row.findall("w:tc", ns)]
        if len(cells) >= 6:
            rows.append((cells[0], float(cells[4]), float(cells[5])))
    strains = pd.DataFrame(rows, columns=["strain", "female", "male"])
    itp = pd.read_csv(root / "data/derived/sex_survival_context.csv")
    itp = itp[itp.cohort_set.eq("all_17_cohorts")].set_index("sex").km_median_days
    values = {"inbred_strains": len(strains),
              "inbred_median_female_minus_male_days": (strains.female - strains.male).median(),
              "inbred_strains_above_itp_female_median": int((strains.female > itp["female"]).sum())}
    assert values["inbred_strains"] == 31
    assert values["inbred_median_female_minus_male_days"] == 9
    assert values["inbred_strains_above_itp_female_median"] == 2
    path = root / "data/derived/strain_context_summary.csv"
    existing = pd.read_csv(path)
    pd.concat([existing, pd.DataFrame(values.items(), columns=["quantity", "value"])],
              ignore_index=True).to_csv(path, index=False)


def main():
    summarize_comparator_survival()
    summarize_weight_lifespan_correlations()
    summarize_sex_survival()
    summarize_founder_context()
    summarize_control_lifespan()
    verify_strain_survival_summary()

if __name__ == "__main__":
    main()
