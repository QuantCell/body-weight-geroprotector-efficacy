"""Verify regenerated results against the manuscript."""
from pathlib import Path
import math
import pandas as pd
import numpy as np
from data_utils import clean_adult_weights, initiation_months
from make_tables import build as verify_supplementary_tables, export_result_tables
root = Path(__file__).absolute().parents[2]
out = root / "data" / "derived"
inputs = root / "data" / "inputs"
verify_supplementary_tables(check=True)
export_result_tables(check=True)
# Excel compatibility prefixes must remain declared after worksheet updates.
from zipfile import ZipFile
from io import BytesIO
from xml.etree import ElementTree as ET
with ZipFile(root / "supplementary/Supplementary_Data_1-10.xlsx") as package:
    for name in package.namelist():
        if not name.endswith(".xml"):
            continue
        raw = package.read(name)
        assert b"absPath" not in raw, f"Workbook contains a local path: {name}"
        if name == "docProps/core.xml":
            assert all(element.text in (None, "") for element in ET.fromstring(raw)
                       if element.tag.endswith("}lastModifiedBy")), "Workbook author field is not empty"
        assert b"16-hydroxyestriol" not in raw, f"Outdated compound label: {name}"
        declared = {prefix for _, (prefix, uri) in ET.iterparse(BytesIO(raw), events=["start-ns"])}
        for element in ET.fromstring(raw).iter():
            prefixes = element.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "").split()
            assert set(prefixes) <= declared, f"Undeclared Excel namespace: {name}"
inventory = pd.read_excel(root / "supplementary/Supplementary_Data_1-10.xlsx", sheet_name="SD1_ITP_experiments")
ursolic = inventory.loc[inventory.cohort.eq("C2013") & inventory.intervention.eq("UA")]
assert len(ursolic) == 2 and ursolic.publication_pmid.astype(str).eq("30688027").all()
udca = inventory.loc[inventory.cohort.eq("C2011") & inventory.intervention.eq("UDCA")]
assert len(udca) == 2 and udca.publication_pmid.astype(str).eq("27312235").all()
control_source = pd.read_excel(root / "supplementary/Supplementary_Data_1-10.xlsx", sheet_name="SD10_control_lifespan")
assert "itp_status" in control_source.columns
olfactory = pd.read_csv(out / "olfactory_cohort_mice.csv")
assert len(olfactory) == 360 and olfactory.UMN.is_unique
family = pd.read_csv(out / "olfactory_cohort_interactions.csv")
assert len(family) == 12
assert family.groupby(["weight_age", "sex", "exposure"]).size().eq(1).all()
assert "variance_cluster" not in family.columns
primary_family = pd.read_excel(root / "supplementary/Supplementary_Data_1-10.xlsx", sheet_name="SD8_Garratt_interactions")
assert len(primary_family) == 12
pd.testing.assert_frame_equal(
    primary_family,
    family.assign(
        weight_age=family.weight_age.map({"12weeks": "12 weeks", "6months": "6 months", "12months": "12 months"}),
        sex=family.sex.map({"F": "Female", "M": "Male"}),
        exposure=family.exposure.map({"FU": "Adult-female odour", "MU": "Adult-male odour"})),
    check_exact=False, rtol=1e-10, atol=1e-12,
)
assert set(family.weight_age) == {"12weeks", "6months", "12months"}
assert (family.lo < family.HR).all() and (family.HR < family.hi).all()
assert family.p.between(0, 1).all()
gradients = pd.read_csv(out / "olfactory_cohort_weight_mortality.csv")
female_control = gradients.query("analysis == 'control_6m' and sex == 'F'").iloc[0]
assert int(female_control.n) == 59
assert math.isclose(female_control.HR_per_SD, 1.51834918738642, abs_tol=1e-10)
assert round(female_control.p, 6) == 0.004204
source = pd.read_excel(inputs / "ITP_bodyweight_all_cohorts.xlsx", sheet_name=0)
adult_columns = ["bw_6", "bw_12", "bw_18", "bw_24"]
withheld = source.qc_note.fillna("").str.contains("metadata conflict", regex=False)
assert withheld.sum() == 2 and source.loc[withheld, adult_columns].isna().all().all()
assert set(source.loc[withheld, "id"]) == {"UM80437", "UM80444"}
confirmed = source.id.isin([
    "UT04088", "UM33789", "UM33796", "JL07798", "JL07799", "JL07800",
    "JL07801", "JL07802", "JL07803", "UM108752", "UM105945",
])
assert confirmed.sum() == 11 and source.loc[confirmed, adult_columns].notna().sum().sum() == 40
pd.testing.assert_frame_equal(
    source, pd.read_excel(root / "supplementary/Supplementary_Data_1-10.xlsx",
                          sheet_name="SD2_mouse_level"), check_dtype=False,
)
pbn = source.query("cohort == 'C2004' and group == '4-OH-PBN'")
assert len(pbn) == 324 and pbn.bw_6.notna().sum() == 277
assert pbn.loc[pbn.site.eq("TJL"), "bw_6"].notna().sum() == 106
predosing = pd.read_csv(out / "predosing_comparison_mice.csv")
assert predosing.body_weight.between(10, 100).all()
withheld_keys = set(source.loc[withheld, "cohort"] + "|" + source.loc[withheld, "id"])
assert not set(predosing.mouse).intersection(withheld_keys)
assert not set(pd.read_csv(out / "primary_model_mice.csv").mouse).intersection(withheld_keys)
boundary = pd.DataFrame({"bw_6": [9.9, 10, 80, 90, 100, 100.1, 361]})
assert clean_adult_weights(boundary).bw_6.isna().tolist() == [True, False, False, False, False, True, True]
assert boundary.bw_6.iloc[-1] == 361
qc = pd.read_csv(out / "adult_weight_qc.csv")
assert int(qc.excluded.sum()) == 9 and int((~qc.excluded).sum()) == 9
assert qc.loc[qc.id.eq("JL23127"), "analysis_weight_g"].iloc[0] == 90
assert qc.loc[qc.id.eq("JL07624"), "analysis_weight_g"].isna().all()
prepared_gain = pd.read_csv(out / "adult_weight_change_mice.csv")
assert len(prepared_gain) == 32277 and prepared_gain.bw_12.eq(90).sum() == 1
assert prepared_gain[["bw_6", "bw_12"]].ge(10).all().all()
assert prepared_gain[["bw_6", "bw_12"]].le(100).all().all()
derived = {
    "weight_effects": pd.read_csv(out / "experiment_weight_effects.csv"),
    "interactions": pd.read_csv(out / "experiment_interactions.csv"),
    "comparator_cells": pd.read_csv(out / "comparator_effect_cells.csv"),
}
assert len(derived["weight_effects"]) == 368
pbn_effects = derived["weight_effects"].query("cohort == 'C2004' and group == '4-OH-PBN' and post_dosing")
assert (pbn_effects.diff_g * pbn_effects.n_tx).sum() / pbn_effects.n_tx.sum() > -1
assert len(derived["interactions"]) == 180
assert len(derived["comparator_cells"]) == 534
assert len(derived["comparator_cells"].query("cohort == 'C2021' and site == 'TJL'")) == 12
# The formulation remains in broader comparisons, not single-compound summaries.
for frame in [*derived.values(), prepared_gain]:
    assert (frame.cohort.eq("C2011") & frame.group.eq("Prot")).any()
for filename, column in [("primary_model_mice.csv", "experiment_sex"),
                         ("predosing_comparison_mice.csv", "experiment"),
                         ("landmark_model_mice.csv", "experiment")]:
    frame = pd.read_csv(out / filename)
    assert frame[column].str.contains("C2011|Prot", regex=False).any()
primary_rows = pd.read_csv(out / "primary_model_mice.csv")
experiment_meta = pd.read_csv(out / "experiment_interactions.csv")
assert experiment_meta.weight_age.eq(6).all()
starts = source.assign(start=initiation_months(source.age_initiation))
starts = starts[~starts.group.eq("Control")].groupby(["cohort", "group"]).start.first()
for experiment, rows in primary_rows.groupby("experiment_sex"):
    cohort, group, sex = experiment.split("|")
    expected_entry = max(6.0, float(starts.loc[(cohort, group)])) * 30.4
    assert rows.entry.sub(expected_entry).abs().lt(1e-9).all()
primary = pd.read_csv(out / "primary_moderation.csv").set_index("sex")
assert int(primary.loc["m", "experiments"]) == 92 and int(primary.loc["m", "unique_mice"]) == 17943
assert int(primary.loc["f", "experiments"]) == 88 and int(primary.loc["f", "unique_mice"]) == 15837
assert math.isclose(primary.loc["m", "HR"], 0.944851051866112, rel_tol=0, abs_tol=1e-8)
split = pd.read_csv(out / "weight_lowering_split.csv").set_index("analysis")
assert int(split.loc["male_weight_lowering", "experiments"]) == 42
assert math.isclose(split.loc["male_weight_lowering", "HR"], 0.896834917, abs_tol=1e-6)
assert int(split.loc["male_predosing_weight_lowering", "experiments"]) == 30
nonlowering = split.loc["male_predosing_nonlowering"]
assert int(nonlowering.experiments) == 30 and round(nonlowering.HR, 2) == 1.01
assert round(nonlowering.lo, 2) == 0.96 and round(nonlowering.hi, 2) == 1.07
assert round(nonlowering.p, 2) == 0.64
q = pd.read_csv(out / "matched_quartile_standardized_survival.csv")
assert q["gain_pct"].round(1).tolist() == [6.2, 8.0, 12.3, 20.2]
sensitivity = pd.read_csv(out / "descriptive_sensitivities.csv")
quarters = sensitivity[sensitivity.analysis.eq("weight-lowering standardized quarters")].sort_values("bin")
for actual, expected in [
    ("control_median", "standardized_control_median"),
    ("treated_median", "standardized_treated_median"),
    ("mean_control_weight", "mean_control_weight"), ("gain_pct", "gain_pct"),
]:
    assert all(math.isclose(a, b, rel_tol=1e-10, abs_tol=1e-10)
               for a, b in zip(quarters[actual], q[expected]))
assert len(quarters) == 4
for label, bins, lowest, highest in [("thirds", 3, 7.3, 20.4), ("fifths", 5, 4.0, 22.2)]:
    rows = sensitivity[sensitivity.analysis.eq("weight-lowering standardized " + label)].sort_values("bin")
    assert len(rows) == bins and rows.bin.tolist() == list(range(1, bins + 1))
    assert round(rows.iloc[0].gain_pct, 1) == lowest
    assert round(rows.iloc[-1].gain_pct, 1) == highest
lightest_fifth = sensitivity[sensitivity.analysis.eq("weight-lowering standardized fifths") & sensitivity.bin.eq(1)].iloc[0]
assert lightest_fifth.control_median == 907 and lightest_fifth.treated_median == 943
assert lightest_fifth.treated_median - lightest_fifth.control_median == 36
assert math.isclose(lightest_fifth.gain_pct, 100 * 36 / 907, rel_tol=1e-12)
conventional = sensitivity[sensitivity.analysis.eq("weight-lowering conventional quartiles")]
assert conventional.gain_pct.round(1).tolist() == [4.5, 10.7, 11.5, 20.0]
trend = pd.read_csv(out / "matched_quartile_trend.csv").iloc[0]
assert int(trend.experiments) == 30 and int(trend.unique_mice) == 8137
assert math.isclose(trend.HR_per_quartile, 0.908736572154869, rel_tol=0, abs_tol=1e-8)
nonlower = pd.read_csv(out / "nonlowering_comparison_mice.csv")
treated_nonlower = nonlower[nonlower["Tx"].eq(1)]
assert len(treated_nonlower[["cohort", "group"]].drop_duplicates()) == 30
assert nonlower["mouse"].nunique() == 8252
comp = pd.read_csv(out / "compound_descriptive.csv")
assert len(comp) == 15 and int(comp.iloc[-1].experiments) == 26 and int(comp.iloc[-1].unique_mice) == 7238
assert "Protandim" not in comp.compound.tolist()
assert round(comp.iloc[-1].lowest_gain_pct,1) == 5.0 and round(comp.iloc[-1].highest_gain_pct,1) == 20.2
assert round(comp.iloc[-1].overall_gain_pct, 1) == 9.7
assert int((comp.iloc[:-1].highest_gain_pct > comp.iloc[:-1].lowest_gain_pct).sum()) == 13
named = comp.set_index("compound")
for compound, lowest, highest in [
    ("Acarbose", 7.9, 24.8), ("17alpha-estradiol", 5.9, 22.9),
    ("Canagliflozin", 6.1, 23.2), ("Rapamycin", 7.3, 21.0),
]:
    assert round(named.loc[compound, "lowest_gain_pct"], 1) == lowest
    assert round(named.loc[compound, "highest_gain_pct"], 1) == highest
common = pd.read_csv(out / "compound_common_interaction.csv").iloc[0]
assert int(common.compounds) == 14 and round(common.estimate, 2) == 0.9
assert int(common.experiments) == 26 and int(common.unique_mice) == 7238
assert round(common.p, 4) == 0.0001
comparator = pd.read_csv(out / "comparator_state.csv")
tertiles = comparator[comparator.split.eq(3)].set_index(["sex", "bin"])
assert round(tertiles.loc[("m", "Q1"), "mean_cell_km_median_gain_pct"], 1) == 4.0
assert round(tertiles.loc[("m", "Q3"), "mean_cell_km_median_gain_pct"], 1) == 8.3
assert round(tertiles.loc[("f", "Q1"), "mean_cell_km_median_gain_pct"], 1) == 1.9
assert round(tertiles.loc[("f", "Q3"), "mean_cell_km_median_gain_pct"], 1) == 2.8
for sex, bin_, expected in [("m", "Q1", 10.9), ("m", "Q3", 21.5),
                            ("f", "Q1", 5.3), ("f", "Q3", 12.0)]:
    assert round(tertiles.loc[(sex, bin_), "hazard_reduction_pct"], 1) == expected
puberty_input = pd.read_csv(inputs / "itp_gn_puberty_weight_data.csv")
puberty_controls = puberty_input[puberty_input.Tx.eq(0)].drop_duplicates("mouse")
reference = pd.read_csv(out / "puberty_weight_reference.csv").set_index("reference_group")
for key, group in puberty_controls.groupby(["cohort", "site", "sex"]):
    row = reference.loc["|".join(key)]
    assert int(row.controls) == len(group)
    assert math.isclose(row.mean_weight_g, group.weight_42d.mean(), rel_tol=1e-12)
    assert math.isclose(row.sd_weight_g, group.weight_42d.std(), rel_tol=1e-12)
    values = puberty_input[(puberty_input.cohort == key[0]) &
                           (puberty_input.site == key[1]) & (puberty_input.sex == key[2])]
    assert np.allclose((values.weight_42d - row.mean_weight_g) / row.sd_weight_g,
                       values.weight_42d_z, rtol=1e-12, atol=1e-12)
gradient = pd.read_csv(out / "comparator_gradients.csv").set_index(["sex", "analysis"])

age = pd.read_csv(out / "control_weight_gradient_by_age.csv").set_index(["weight_age", "sex"])
expected_age = {
    (6, 'm'): (4983, 1.38918369577528, 1.34762754413104, 1.43202129476527, 7.08523839049341e-100),
    (6, 'f'): (4553, 1.07542292934942, 1.04402639328046, 1.10776363932383, 1.50858332301721e-06),
    (12, 'm'): (4572, 1.36270654257355, 1.3187482442416, 1.40813011830069, 2.11834714458959e-76),
    (12, 'f'): (4537, 1.09870819571753, 1.06621500179951, 1.13219162861101, 7.944224065921061e-10),
    (18, 'm'): (3900, 1.12967041747538, 1.09056219581102, 1.17018108368405, 1.1788650530754e-11),
    (18, 'f'): (4450, 1.05286382919565, 1.02137735283941, 1.08532095385398, 0.0008827113173421),
    (24, 'm'): (2670, 0.874475452887517, 0.837391715426413, 0.913201436813149, 1.3032113417376201e-09),
    (24, 'f'): (3508, 0.969843007572835, 0.936579603987902, 1.00428778860112, 0.0854854448220026),
}
for key, (expected_n, expected_hr, expected_lo, expected_hi, expected_p) in expected_age.items():
    assert int(age.loc[key, "n"]) == expected_n
    for column, expected in (
        ("HR_per_SD", expected_hr), ("lo", expected_lo),
        ("hi", expected_hi), ("p", expected_p),
    ):
        assert math.isclose(age.loc[key, column], expected,
                            rel_tol=1e-6, abs_tol=1e-12)
puberty = pd.read_csv(out / "puberty_moderation.csv").set_index("sex")
assert int(puberty.loc["m", "experiments"]) == 4 and int(puberty.loc["m", "unique_mice"]) == 730
assert round(puberty.loc["m", "HR"], 2) == 1.00 and round(puberty.loc["m", "p"], 2) == 0.97
assert int(puberty.loc["f", "experiments"]) == 2 and int(puberty.loc["f", "unique_mice"]) == 553
assert round(puberty.loc["f", "HR"], 2) == 1.03 and round(puberty.loc["f", "p"], 2) == 0.77
joint = pd.read_csv(out / "puberty_joint.csv").set_index("term")
assert int(joint.loc["w6z", "n"]) == 539 and round(joint.loc["w6z", "HR"], 2) == 1.46
gain = pd.read_csv(out / "early_weight_change.csv")
male_gain = gain[(gain.sex == "m") & (gain.timing == "primary") & gain.adjusted_for_w6].iloc[0]
female_gain = gain[(gain.sex == "f") & (gain.timing == "primary") & gain.adjusted_for_w6].iloc[0]
assert int(male_gain.n_mice) == 14118 and round(male_gain.hr, 2) == 1.1
assert math.isclose(male_gain.p, 3.1220900809913397e-16, rel_tol=1e-6)
assert int(female_gain.n_mice) == 13167 and round(female_gain.hr, 2) == 1.04
assert int(male_gain.cells) == 257 and round(male_gain.ci_low, 2) == 1.07 and round(male_gain.ci_high, 2) == 1.12
assert int(female_gain.cells) == 251 and round(female_gain.ci_low, 2) == 1.03 and round(female_gain.ci_high, 2) == 1.06
maternal = pd.read_csv(out / "maternal_adjustment.csv")
successful = pd.read_csv(out / "successful_landmarks.csv")
male12 = successful[(successful.landmark == 12) & (successful.sex == "m")].iloc[0]
assert int(male12.unique_mice) == 5807 and int(male12.experiments) == 21
assert round(male12.HR, 2) == 0.9 and round(male12.p, 4) == 0.0006
lag = pd.read_csv(out / "successful_landmarks_lag90.csv")
male12_lag = lag[(lag.landmark == 12) & (lag.sex == "m")].iloc[0]
assert round(male12_lag.HR, 2) == 0.91
assert round(male12_lag.lo, 2) == 0.86 and round(male12_lag.hi, 2) == 0.97
assert round(male12_lag.p, 3) == 0.002
loo = pd.read_csv(out / "successful_landmarks_loo12.csv")
male12_loo = loo[(loo.sex == "m") & (loo.drop_cohort != "none")]
assert len(male12_loo) == 10 and (male12_loo.p < 0.05).all()
assert round(male12_loo.HR.min(), 2) == 0.88 and round(male12_loo.HR.max(), 2) == 0.92
strata = pd.read_csv(out / "weight_lifespan_by_experiment.csv")
assert int((strata.loc[strata.sex.eq("f"), "rho"] < 0).sum()) == 87
supplement = root / "supplementary" / "Supplementary_Data_1-10.xlsx"
for sheet, output in [("SD6_experiment_contrasts", "experiment_weight_changes.csv"),
                      ("SD4_stratum_correlations", "weight_lifespan_by_experiment.csv")]:
    pd.testing.assert_frame_equal(pd.read_excel(supplement, sheet_name=sheet),
                                  pd.read_csv(out / output), check_dtype=False,
                                  rtol=1e-10, atol=1e-12)
assert (maternal.HR_weight - maternal.HR_adjusted_maternal).abs().max() < 0.01
maternal_by_sex = maternal.set_index("sex")
for sex, unadjusted, adjusted in [("m", 1.38473940106212, 1.38438675107383),
                                  ("f", 1.1034936669522, 1.10326868920476)]:
    assert math.isclose(maternal_by_sex.loc[sex, "HR_weight"], unadjusted, rel_tol=1e-9)
    assert math.isclose(maternal_by_sex.loc[sex, "HR_adjusted_maternal"], adjusted, rel_tol=1e-9)

from importlib import import_module
import_module("08_weight_response").validate_semaglutide_extraction()
corpus = pd.read_excel(supplement, sheet_name="SD9_cross_study_experiments")
assert len(corpus) == 325 and corpus.pubmed_id.nunique() == 90
cross = pd.read_csv(out / "supplementary_analyses.csv").set_index("analysis")
assert int(cross.loc["pharmacological_sex_by_source", "experiments"]) == 108
expected_slopes = {
    "slope_ITP_Female": (25, 12, 0.46, 0.13, 0.79),
    "slope_ITP_Male": (31, 13, 1.03, 0.72, 1.35),
    "slope_non-ITP pharmacological_Female": (23, 16, 0.92, 0.54, 1.29),
    "slope_non-ITP pharmacological_Male": (29, 23, 0.79, 0.38, 1.20),
    "slope_dietary restriction_Female": (15, 7, 0.79, 0.57, 1.01),
    "slope_dietary restriction_Male": (12, 8, 0.81, 0.65, 0.97),
}
for analysis, (experiments, studies, estimate, ci_low, ci_high) in expected_slopes.items():
    assert int(cross.loc[analysis, "experiments"]) == experiments
    assert int(cross.loc[analysis, "studies"]) == studies
    assert round(cross.loc[analysis, "estimate"], 2) == estimate
    assert round(cross.loc[analysis, "ci_low"], 2) == ci_low
    assert round(cross.loc[analysis, "ci_high"], 2) == ci_high
expected_medians = {
    "median_ratio_ITP_Female": 0.40,
    "median_ratio_ITP_Male": 1.27,
    "median_ratio_dietary restriction_Female": 0.67,
    "median_ratio_dietary restriction_Male": 0.74,
    "median_ratio_non-ITP pharmacological_Female": 1.00,
    "median_ratio_non-ITP pharmacological_Male": 1.29,
}
for analysis, expected in expected_medians.items():
    assert round(cross.loc[analysis, "estimate"], 2) == expected
for analysis, difference, ci_low, ci_high, p_value in [
    ("female_minus_male_ITP", -0.57, -1.02, -0.12, 0.016),
    ("female_minus_male_non-ITP pharmacological", 0.13, -0.32, 0.58, 0.56),
    ("female_minus_male_dietary restriction", -0.02, -0.29, 0.26, 0.90),
]:
    assert round(cross.loc[analysis, "estimate"], 2) == difference
    assert round(cross.loc[analysis, "ci_low"], 2) == ci_low
    assert round(cross.loc[analysis, "ci_high"], 2) == ci_high
    assert float(f"{cross.loc[analysis, 'p_t']:.2g}") == p_value
assert round(cross.loc["pharmacological_sex_by_source", "p_t"], 3) == 0.025
assert round(cross.loc["pharmacological_sex_by_source_weight_loss_ge_5pct", "p_t"], 3) == 0.021
combined = gradient.loc[("m", "within_cohort_shared_control_clustered")]
assert round(combined.hr_per_sd, 2) == 0.96
assert round(combined.ci_low, 2) == 0.93 and round(combined.ci_high, 2) == 1.00
assert round(combined.p, 3) == 0.034
assert round(gradient.loc[("f", "within_cohort_shared_control_clustered"), "p"], 2) == 0.27


control_records = pd.read_excel(root / "supplementary/Supplementary_Data_1-10.xlsx", sheet_name="SD10_control_lifespan")
assert not (control_records.pmid.eq(24245565) & control_records.lifespan_control_days.isin([831, 836])).any()
jones = control_records.loc[control_records.pmid.eq(7140862)].set_index("sex")
assert len(jones) == 2 and jones.loc["Male", "lifespan_control_days"] == 744.1 and jones.loc["Female", "lifespan_control_days"] == 814.8
inoue = control_records.loc[control_records.pmid.eq(12954483)]
assert len(inoue) == 3 and inoue.sex.eq("Male").all() and inoue.lifespan_control_days.eq(672).all()

ceiling = pd.read_csv(out / "control_lifespan_ceiling.csv").set_index(["specification", "sex"])
expected_ceiling = {
    ('unique_control_groups', 'Female'): (18, 891.0, 20, 756.05, 0, 7.188842228610488e-07),
    ('unique_control_groups', 'Male'): (18, 800.5, 24, 793.7, 11, 0.9898589414976676),
    ('study_level_median', 'Female'): (13, 893.0, 18, 733.2, 0, 5.50520338029921e-06),
    ('study_level_median', 'Male'): (14, 795.5, 22, 787.4, 10, 0.7089826427378236),
    ('median_endpoints_only', 'Female'): (18, 891.0, 11, 819.7, 0, 6.913312302260613e-05),
    ('median_endpoints_only', 'Male'): (18, 800.5, 12, 803.9, 6, 0.7507567541261222),
    ('published_2000_or_later', 'Female'): (18, 891.0, 15, 819.7, 0, 6.717906398911111e-06),
    ('published_2000_or_later', 'Male'): (18, 800.5, 18, 807.8, 9, 0.6237627441568994),
}
for key, expected in expected_ceiling.items():
    row = ceiling.loc[key]
    values = (int(row.itp_n), float(row.itp_median_days), int(row.non_itp_n),
              float(row.non_itp_median_days), int(row.non_itp_reaching_itp_median),
              float(row.p_mannwhitney))
    for observed, target in zip(values, expected):
        assert math.isclose(observed, target, rel_tol=1e-6, abs_tol=1e-12)

survival_context = pd.read_csv(out / "sex_survival_context.csv").set_index(
    ["cohort_set", "sex"]
)
full = "all_17_cohorts"
do = survival_context.loc[("DRiDO_ad_libitum", "female")]
assert int(do.n) == 188
assert math.isclose(do.day500_mortality_pct, 100 * 23 / 188, rel_tol=1e-12)
assert int(survival_context.loc[(full, "male"), "n"]) == 5352
assert int(survival_context.loc[(full, "female"), "n"]) == 4833
assert int(survival_context.loc[(full, "male"), "cohorts"]) == 17
assert survival_context.loc[(full, "male"), "km_median_days"] == 800
assert survival_context.loc[(full, "female"), "km_median_days"] == 886
assert survival_context.loc[(full, "female_minus_male"), "km_median_days"] == 86
assert round(survival_context.loc[(full, "male"), "day500_mortality_pct"], 1) == 16.0
assert round(survival_context.loc[(full, "female"), "day500_mortality_pct"], 1) == 2.4
assert round(
    survival_context.loc[(full, "female"), "day500_mortality_all_events_pct"], 1
) == 3.0

six = "published_six_cohorts"
assert int(survival_context.loc[(six, "male"), "cohorts"]) == 6
assert int(survival_context.loc[(six, "male"), "n"]) == 1993
assert int(survival_context.loc[(six, "female"), "n"]) == 1697
assert survival_context.loc[(six, "male"), "km_median_days"] == 803
assert survival_context.loc[(six, "female"), "km_median_days"] == 887
assert survival_context.loc[(six, "female_minus_male"), "km_median_days"] == 84

external_survival = pd.read_csv(out / "external_sex_survival_context.csv").set_index(
    "population"
)
for population, male, female, male_n, female_n, difference in [
    ("Independent UM-HET3", 831, 836, 14, 16, 5),
    ("Independent 5LG6 cohort", 804, 865, 40, 45, 61),
]:
    row = external_survival.loc[population]
    assert int(row.male_survival_days) == male and int(row.female_survival_days) == female
    assert int(row.male_n) == male_n and int(row.female_n) == female_n
    assert int(row.female_minus_male_days) == difference

cohort_survival = pd.read_csv(out / "sex_survival_by_cohort.csv")
assert len(cohort_survival) == 17 and (cohort_survival.female_advantage_days > 0).all()
assert cohort_survival.female_advantage_days.min() == 52
assert cohort_survival.female_advantage_days.max() == 135
cohort_survival = cohort_survival.set_index("cohort")
assert cohort_survival.loc["C2020", "female_km_median_days"] == 906
assert cohort_survival.loc["C2021", "female_km_median_days"] == 870

drido_landmark = pd.read_csv(out / "drido_180_landmark_pooled_cr.csv")
assert len(drido_landmark) == 559 and int(drido_landmark.event.sum()) == 552
assert drido_landmark.groupby("diet").size().to_dict() == {"20": 189, "40": 182, "AL": 188}
assert (drido_landmark["weight_age_days"].sub(180).abs() <= 45).all()
assert set(drido_landmark.groupby("group")["matched_quarter"].nunique()) == {4}

drido_quartiles = pd.read_csv(out / "drido_pooled_cr_quartiles.csv").set_index("quarter")
expected_drido = {
    1: (24.8, 877, 973, 96, 10.9),
    2: (29.8, 825, 974, 149, 18.1),
    3: (34.6, 779, 931, 152, 19.5),
    4: (42.1, 714, 946, 232, 32.5),
}
for quarter, (weight, al, cr, days, percent) in expected_drido.items():
    row = drido_quartiles.loc[quarter]
    assert round(row.mean_weight_g, 1) == weight
    assert float(row.al_km_median_days) == al and float(row.cr_km_median_days) == cr
    assert float(row.gain_days) == days and round(row.gain_pct, 1) == percent
drido_interaction = pd.read_csv(out / "drido_pooled_cr_interaction.csv").iloc[0]
assert int(drido_interaction["n"]) == 559 and int(drido_interaction.events) == 552
assert round(drido_interaction.hr, 2) == 0.90
assert round(drido_interaction.ci_low, 2) == 0.76
assert round(drido_interaction.ci_high, 2) == 1.05
assert round(drido_interaction.p, 2) == 0.18
drido_gradients = pd.read_csv(out / "drido_weight_mortality_by_group.csv").set_index("group")
assert round(drido_gradients.loc["AL", "hr"], 2) == 1.21
assert round(drido_gradients.loc["AL", "p"], 3) == 0.015
assert round(drido_gradients.loc["CR", "hr"], 2) == 1.05

strain = pd.read_csv(out / "strain_context_summary.csv").set_index("quantity")
assert int(strain.loc["age_matched_strains", "value"]) == 27
assert round(float(strain.loc["survey_mean_weight_g", "value"]), 1) == 30.3
for quantity, expected in {
    "weight_survey_strains": 27,
    "weight_strains_at_or_above_itp_mean": 1,
    "inbred_strains": 31,
    "inbred_median_female_minus_male_days": 9,
    "inbred_strains_above_itp_female_median": 2,
    "weight_strains_at_or_above_lightest_quarter": 7,
    "male_fat_strains": 40,
    "male_fat_at_or_above_itp": 2,
    "female_fat_strains": 40,
    "female_fat_at_or_above_itp": 1,
}.items():
    assert int(strain.loc[quantity, "value"]) == expected
founder = pd.read_csv(out / "founder_strain_weight_comparison.csv").set_index("quantity")
for quantity, expected in {
    "four_founder_mean_weight_g": 29.4,
    "itp_male_control_6m_n": 4983,
    "itp_male_control_6m_mean_weight_g": 40.6,
    "itp_percent_above_founder_mean": 38.0,
    "lowest_itp_quarter_mean_weight_g": 34.4,
    "heaviest_founder_reported_mean_weight_g": 32.0,
}.items():
    assert round(float(founder.loc[quantity, "value"]), 1) == expected
assert int(founder.loc["itp_male_control_6m_n", "value"]) == int(age.loc[(6, "m"), "n"])
assert round(drido_gradients.loc["CR", "ci_low"], 2) == 0.95
assert round(drido_gradients.loc["CR", "ci_high"], 2) == 1.17

# Treated-only sensitivity and the overlapping trajectory populations.
treated = gain[(gain.sex == "m") & gain.timing.eq("treated_primary") & gain.adjusted_for_w6].iloc[0]
assert int(treated.n_mice) == 9705 and int(treated.cells) == 208
assert round(treated.hr, 2) == 1.1
assert round(treated.ci_low, 2) == 1.07 and round(treated.ci_high, 2) == 1.13
assert math.isclose(treated.p, 5.422056128688749e-11, rel_tol=1e-6)
trajectory = prepared_gain[prepared_gain.sex.eq("m")].copy()
trajectory = trajectory[trajectory.groupby("cell").dead.transform("sum") >= 12]
controls = trajectory[trajectory.is_control]
early_treated = trajectory[~trajectory.is_control & trajectory.init_months.lt(12)]
late_treated = trajectory[~trajectory.is_control & trajectory.init_months.ge(12)]
assert (len(controls), controls.cell.nunique()) == (4413, 49)
assert (len(early_treated), early_treated.cell.nunique()) == (9705, 208)
assert (len(late_treated), late_treated.cell.nunique()) == (2697, 58)

# Check weight filtering and experiment eligibility.
experiments = pd.read_csv(out / "experiment_weight_changes.csv")
assert len(experiments) == 163 and experiments.groupby("sex").size().to_dict() == {'f': 80, 'm': 83}
from data_utils import initiation_months
assert initiation_months(pd.Series(["10, then 17", "6", 12, None])).iloc[:3].tolist() == [10, 6, 12]
prot = experiments[experiments.cohort.eq("C2011") & experiments.group.eq("Prot")]
assert len(prot) == 2 and prot.age_init.eq(10).all()
assert experiments.age_init.lt(18).all()
assert experiments[["n_tw", "n_cw", "n_td", "n_cd"]].ge(8).all().all()
enal = experiments[(experiments.cohort == "C2005") & (experiments.sex == "m") & (experiments.group == "Enal")].iloc[0]
assert int(enal.n_tw) == 132, "C2005 male Enal weight eligibility"
assert math.isclose(enal.wt_pct, 0.6550314547721157, abs_tol=1e-10)

quarters = pd.read_csv(out / "weight_change_quartiles.csv").set_index(["sex", "quartile"])
for sex, sizes, means in [("m", [24, 16, 15, 19], [10.9, 4.1, 4.4, 2.7]),
                          ("f", [22, 15, 14, 17], [4.6, 1.4, 1.5, 0.6])]:
    assert quarters.loc[sex, "experiments"].tolist() == sizes
    assert quarters.loc[sex, "mean_lifespan_change_pct"].round(1).tolist() == means
trends = pd.read_csv(out / "weight_change_trend.csv").set_index("sex")
assert (int(trends.loc["m", "n"]), int(trends.loc["m", "clusters"])) == (74, 13)
assert (int(trends.loc["f", "n"]), int(trends.loc["f", "clusters"])) == (68, 12)
assert round(trends.loc["m", "estimate"], 2) == -2.56
assert round(trends.loc["m", "lo"], 2) == -4.6 and round(trends.loc["m", "hi"], 2) == -0.53
assert round(trends.loc["m", "p"], 4) == 0.0178
assert round(trends.loc["f", "p"], 3) == 0.098
slopes = pd.read_csv(out / "weight_change_slopes.csv").set_index(["sex", "specification"])
assert round(slopes.loc[("m", "pooled"), "estimate"], 2) == -0.71
assert round(slopes.loc[("m", "pooled"), "p"], 4) == 0.0016
assert math.isclose(slopes.loc[("m", "pooled"), "se"], 0.1836229411527838,
                    rel_tol=1e-9), "Weight-response finite-sample variance correction"
assert round(slopes.loc[("m", "within_cohort"), "estimate"], 2) == -0.65
assert round(slopes.loc[("m", "within_cohort"), "p"], 3) == 0.008
assert round(slopes.loc[("f", "within_cohort"), "p"], 2) == 0.28
joint = pd.read_csv(out / "weight_change_vs_comparator.csv").set_index(["sex", "model", "term"])
male = joint.loc[("m", "both_together", "wt_pct")]
assert int(male.n) == 246 and int(male.clusters) == 48
assert round(male.estimate, 4) == 0.0214
assert round(male.lo, 4) == 0.0137 and round(male.hi, 4) == 0.0291
assert float(f"{male.p:.1e}") == 1.2e-06
assert round(joint.loc[("m", "both_together", "ctl_bw6"), "p"], 2) == 0.05
assert round(joint.loc[("f", "both_together", "wt_pct"), "p"], 3) == 0.004

# Verify the specified t-reference distribution and intervals.
from scipy.stats import t as student_t
for filename in ("weight_change_trend.csv", "weight_change_slopes.csv", "weight_change_vs_comparator.csv"):
    results = pd.read_csv(out / filename).dropna(subset=["se", "df"])
    for row in results.itertuples(index=False):
        assert row.df == row.clusters - 1
        expected_p = 2 * student_t.sf(abs(row.estimate / row.se), row.df)
        assert math.isclose(row.p, expected_p, rel_tol=1e-9, abs_tol=1e-14), filename + ": t reference"
        critical = student_t.ppf(0.975, row.df)
        assert math.isclose(row.lo, row.estimate - critical * row.se, rel_tol=1e-9, abs_tol=1e-12)
        assert math.isclose(row.hi, row.estimate + critical * row.se, rel_tol=1e-9, abs_tol=1e-12)

from openpyxl import load_workbook
supplement = load_workbook(root / "supplementary/Supplementary_Data_1-10.xlsx", read_only=True, data_only=True)
assert supplement["SD6_experiment_contrasts"].max_row == 164
control_experiments = list(supplement["SD10_control_lifespan"].values)
assert len(control_experiments) == 231
assert not any("doxorubicin" in str(row[-1]).lower() for row in control_experiments[1:])
matched = supplement["SD3_matched_quartiles"]
assert matched.cell(11, 2).value == 30 and matched.cell(11, 4).value == 8137
assert matched.cell(5, 5).value == 890 and matched.cell(8, 6).value == 809
assert math.isclose(matched.cell(5, 8).value, 100 * 55 / 890, abs_tol=1e-10)
supplement.close()

rmst = pd.read_csv(out / "quartile_rmst.csv")
assert len(rmst) == 8 and (rmst.horizon_age_days == 1000).all()
assert (rmst.entry_age_days == 182.625).all() and (rmst.experiments == 30).all()
assert (rmst.unique_mice == 8137).all()
assert all(math.isclose(r.gain_days, r.treated_rmst_days - r.control_rmst_days,
                        abs_tol=1e-9) for r in rmst.itertuples())
ranked = rmst[rmst.grouping == "separate_ranks"].set_index("quartile")
assert round(ranked.loc[1, "gain_days"], 1) == 37.3
assert round(ranked.loc[4, "gain_days"], 1) == 88.9
contrasts = pd.read_csv(out / "quartile_rmst_contrasts.csv").set_index("grouping")
assert (contrasts.bootstrap_replicates == 2000).all()
assert round(contrasts.loc["separate_ranks", "heaviest_minus_lightest_gain_days"], 1) == 51.6
assert round(contrasts.loc["separate_ranks", "lo"], 1) == 26.7
assert round(contrasts.loc["separate_ranks", "hi"], 1) == 76.9
assert round(contrasts.loc["common_control_thresholds", "heaviest_minus_lightest_gain_days"], 1) == 52.4
common = pd.read_csv(out / "common_threshold_quartile_trend.csv").iloc[0]
assert int(common.sites_by_experiment) == 88 and int(common.comparison_rows) == 12942
assert round(common.HR_per_quartile, 2) == 0.92
assert round(common.lo, 2) == 0.87 and round(common.hi, 2) == 0.96
assert round(common.p, 4) == 0.0002
thresholds = pd.read_csv(out / "common_weight_thresholds.csv")
assert len(thresholds) == 89 and thresholds.retained.sum() == 88
assert (thresholds.q25_g < thresholds.q50_g).all()
assert (thresholds.q50_g < thresholds.q75_g).all()

site_estimates = pd.read_csv(out / "site_moderation.csv")
assert len(site_estimates) == 6 and set(site_estimates.sex) == {"f", "m"}
for sex in ["f", "m"]:
    assert set(site_estimates.loc[site_estimates.sex.eq(sex), "site"]) == {"TJL", "UM", "UT"}
assert site_estimates.loc[site_estimates.sex.eq("f"), "HR"].round(2).between(1.01, 1.05).all()
male_sites = site_estimates.loc[site_estimates.sex.eq("m"), "HR"]
assert set(male_sites.round(2)) == {0.94, 0.95}
site_tests = pd.read_csv(out / "site_heterogeneity.csv").set_index("sex")
assert site_tests.df.eq(2).all()
assert round(site_tests.loc["f", "p"], 2) == 0.62
assert round(site_tests.loc["m", "p"], 2) == 0.99
assert (root / "figures/Figure_4_compound_weight_strata.png").is_file()

print("Verified the principal, class, quartile, compound, comparator, developmental, cross-study, "
      "treated-only trajectory and QC-cleaned weight-response results.")

# Founder medians are read from the published supplementary table.
import importlib.util
figure_spec = importlib.util.spec_from_file_location("figure_sources", root / "scripts/analysis/make_figures.py")
figure_sources = importlib.util.module_from_spec(figure_spec)
figure_spec.loader.exec_module(figure_sources)
founder_survival = figure_sources.load_founder_lifespans(root).set_index(["Strain", "Sex"])
for strain, male, female in [("BALB/cByJ", 710.5, 790), ("C57BL/6J", 901, 914),
                            ("C3H/HeJ", 721, 659), ("DBA/2J", 697.5, 683.5)]:
    assert float(founder_survival.loc[(strain, "M"), "median"]) == male
    assert float(founder_survival.loc[(strain, "F"), "median"]) == female

# Check survival curves against standardized quartile medians.
curves = pd.read_csv(out / "matched_quartile_survival_curves.csv")
quartile_medians = pd.read_csv(out / "matched_quartile_standardized_survival.csv").set_index("quartile")
for quarter in range(1, 5):
    for treatment, column in [(0, "standardized_control_median"),
                               (1, "standardized_treated_median")]:
        curve = curves[(curves.quartile == quarter) & (curves.Tx == treatment)]
        assert curve.time.is_monotonic_increasing
        assert curve.survival.between(0, 1).all()
        assert curve.survival.diff().dropna().le(1e-12).all()
        median = curve.loc[curve.survival.le(0.5), "time"].iloc[0]
        assert median == quartile_medians.loc[quarter, column]

rapamycin_curves = pd.read_csv(out / "rapamycin_quartile_survival_curves.csv")
rapamycin_summary = pd.read_csv(out / "compound_descriptive.csv").set_index("compound").loc["Rapamycin"]
for quarter, prefix in [(1, "lowest"), (4, "highest")]:
    for tx, group in [(0, "control"), (1, "treated")]:
        curve = rapamycin_curves[(rapamycin_curves.quartile == quarter) & (rapamycin_curves.Tx == tx)]
        assert curve.loc[curve.survival.le(0.5), "time"].iloc[0] == rapamycin_summary[f"{prefix}_{group}_median"]
assert (root / "figures/Extended_Data_Figure_1_all_experiment_interactions.png").is_file()
