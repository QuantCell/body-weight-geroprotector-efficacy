"""Build the external-dataset assessment and ITP experiment inventory."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile
import argparse
import csv
import io
import math
import re
import zipfile

from openpyxl import load_workbook
import numpy as np
import pandas as pd


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "derived"
SUP = ROOT / "supplementary"


def xml_bytes(root):
    """Serialize worksheet XML with its Excel compatibility namespaces."""
    namespaces = {
        "x14ac": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac",
        "xr": "http://schemas.microsoft.com/office/spreadsheetml/2014/revision",
        "xr2": "http://schemas.microsoft.com/office/spreadsheetml/2015/revision2",
        "xr3": "http://schemas.microsoft.com/office/spreadsheetml/2016/revision3",
        "x16r2": "http://schemas.microsoft.com/office/spreadsheetml/2015/02/main",
    }
    raw = ET.tostring(root, encoding="utf-8", xml_declaration=True).decode("utf-8")
    start = re.search(r"<(?!\?)[\w:]+\b[^>]*>", raw)
    declared = set(re.findall(r"xmlns:([\w]+)=", start.group()))
    needed = set()
    for element in root.iter():
        needed.update(element.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "").split())
    missing = needed - declared
    extra = "".join(f' xmlns:{prefix}="{namespaces[prefix]}"' for prefix in sorted(missing))
    return (raw[:start.end() - 1] + extra + raw[start.end() - 1:]).encode("utf-8")


def joined(values) -> str:
    return "; ".join(sorted({str(v).strip() for v in values if pd.notna(v)}))


def experiment_inventory() -> pd.DataFrame:
    mice = pd.read_excel(ROOT / "data/inputs/ITP_bodyweight_all_cohorts.xlsx")
    primary = pd.read_csv(OUT / "primary_model_mice.csv")
    meta = pd.read_csv(OUT / "experiment_interactions.csv")
    effects = pd.read_csv(OUT / "experiment_weight_effects.csv")
    quartiles = pd.read_csv(OUT / "predosing_comparison_mice.csv")
    references = pd.read_csv(ROOT / "data/inputs/itp_experiment_references.csv").fillna("")
    assert not references.duplicated(["cohort", "group"]).any()
    publications = {(r["cohort"], r["group"]): r for r in references.to_dict("records")}
    assert not mice.duplicated(["cohort", "id"]).any()
    assert not primary.duplicated(["experiment_sex", "mouse"]).any()
    keys = ["cohort", "group", "sex"]
    metadata = {tuple(r[k] for k in keys): r for r in meta.to_dict("records")}
    primary_groups = {k: g for k, g in primary.groupby("experiment_sex", sort=True)}
    post = effects[effects.post_dosing.astype(str).str.lower().eq("true")].copy()
    post["weighted_diff"] = post.diff_g * post.n_tx
    averages = post.groupby(["cohort", "group"])[["weighted_diff", "n_tx"]].sum()
    lowering = (averages.weighted_diff / averages.n_tx).lt(-1).to_dict()
    quartile_keys = set(quartiles.experiment)
    controls = {(c, s): g for (c, s), g in mice[mice.group.eq("Control")].groupby(["cohort", "sex"])}
    treatment = mice[~mice.group.eq("Control")]
    comparisons = treatment.groupby(["cohort", "sex"]).group.nunique().to_dict()
    rows = []
    for (cohort, group, sex), raw in treatment.groupby(keys, sort=True):
        key = f"{cohort}|{group}|{sex}"
        ctl = controls[(cohort, sex)]
        d = primary_groups.get(key, primary.iloc[:0])
        tx, co = d[d.Tx.eq(1)], d[d.Tx.eq(0)]
        schedules = joined(raw.age_initiation)
        match = re.search(r"\d+(?:\.\d+)?", schedules)
        reference = metadata.get((cohort, group, sex), {})
        sites = joined(raw.site)
        primary_sites = joined(d.stratum.str.split("|", regex=False).str[0])
        control_sites = sorted(set(raw.site) & set(ctl.site))
        publication = publications.get((cohort, group), {})
        if len(d):
            assert publication.get("reference"), f"Missing publication: {key}"
        rows.append({
            "experiment_id": key,
            "cohort": cohort,
            "intervention": group,
            "intervention_full_name": "Ursolic acid" if group == "UA" else publication.get("intervention", group),
            "sex": {"m": "Male", "f": "Female"}[sex],
            "recorded_dose": joined(raw.dose),
            "recorded_start_schedule_months": schedules,
            "first_start_months": float(match.group()) if match else None,
            "source_sites": sites,
            "treated_mice_in_source": raw.id.nunique(),
            "control_mice_in_source": ctl.id.nunique(),
            "primary_model": "Yes" if len(d) else "No",
            "primary_exclusion_reason": "" if len(d) else "No recorded six-month weight",
            "primary_weight_age_months": reference.get("weight_age") if len(d) else None,
            "follow_up_start_months": float(d.entry.iloc[0]) / 30.4 if len(d) else None,
            "primary_weight_before_dosing": "Yes" if len(d) and match and float(reference["weight_age"]) < float(match.group()) else "No",
            "primary_sites": primary_sites,
            "primary_treated_mice": tx.mouse.nunique(),
            "primary_control_mice": co.mouse.nunique(),
            "primary_treated_deaths": int(tx.dead.sum()),
            "primary_control_deaths": int(co.dead.sum()),
            "weight_lowering_class": ("Weight-lowering" if lowering.get((cohort, group), False) else "Non-lowering") if len(d) else "Not in primary model",
            "male_predosing_quartile_analysis": "Yes" if sex == "m" and f"{cohort}|{group}" in quartile_keys else "No",
            "shared_control_source_groups": "; ".join(f"{cohort}|{site}|{sex}|Control" for site in control_sites),
            "other_source_comparisons_sharing_controls": comparisons[(cohort, sex)] - 1,
            "source_url": "https://phenome.jax.org/centers/ITP",
            "itp_publication": publication.get("reference", ""),
            "publication_pmid": publication.get("pmid", ""),
            "publication_doi": publication.get("doi", ""),
            "mpd_project": publication.get("mpd_project", "ITP" + str(cohort).removeprefix("C")),
        })
    result = pd.DataFrame(rows).sort_values("primary_model", ascending=False, kind="stable").reset_index(drop=True)
    included = set(result.loc[result.primary_model.eq("Yes"), "experiment_id"])
    assert included == set(primary.experiment_sex)
    for sex in primary.sex.unique():
        selected = primary[primary.sex.eq(sex)]
        benchmark = pd.read_csv(OUT / "primary_moderation.csv")
        expected = benchmark[benchmark.sex.eq(sex)].iloc[0]
        assert selected.experiment_sex.nunique() == int(expected["experiments"])
        assert selected.mouse.nunique() == int(expected["unique_mice"])
    for r in result[result.primary_model.eq("Yes")].to_dict("records"):
        key = (r["cohort"], r["intervention"], {"Male": "m", "Female": "f"}[r["sex"]])
        assert r["primary_treated_mice"] + r["primary_control_mice"] == int(metadata[key]["n"])
    return result.drop(columns=["shared_control_source_groups", "publication_doi", "source_url"]).rename(
        columns=lambda name: name.removeprefix("primary_"))


def csv_text(frame: pd.DataFrame) -> str:
    return frame.to_csv(index=False, lineterminator="\n", float_format="%.12g")


def sync_excel_tables(products, check):
    """Update the supplementary workbook from CSV tables."""
    path = SUP / "Supplementary_Data_1-10.xlsx"
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    tag = lambda name: "{" + ns + "}" + name
    with ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    shared_strings = []
    if "xl/sharedStrings.xml" in parts:
        shared_strings = ["".join(node.itertext())
                          for node in ET.fromstring(parts["xl/sharedStrings.xml"])]
    styles = ET.fromstring(parts["xl/styles.xml"])
    cell_styles = styles.find(tag("cellXfs"))
    no_wrap_styles = {}
    relationships = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    targets = {r.get("Id"): r.get("Target") for r in relationships}
    book = ET.fromstring(parts["xl/workbook.xml"])
    sheets = {}
    for sheet in book.findall(tag("sheets") + "/" + tag("sheet")):
        target = targets[sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")]
        sheets[sheet.get("name")] = target.lstrip("/") if target.startswith("/") else "xl/" + target
    numeric = {"follow_up_start_months", "first_start_months", "treated_mice_in_source", "control_mice_in_source", "weight_age_months", "treated_mice", "control_mice", "treated_deaths", "control_deaths", "other_source_comparisons_sharing_controls", "mice", "deaths", "clusters", "HR", "lo", "hi", "p"}
    for sheet_name, content in products.items():
        sheet = sheets[sheet_name]
        root = ET.fromstring(parts[sheet])
        rows = list(csv.reader(io.StringIO(content)))
        xml_rows = root.find(tag("sheetData")).findall(tag("row"))
        assert len(rows) == len(xml_rows), sheet_name
        for ri, (values, row) in enumerate(zip(rows, xml_rows)):
            cells = {c.get("r"): c for c in row.findall(tag("c"))}
            for ci, value in enumerate(values):
                col, k = "", ci + 1
                while k:
                    k, rem = divmod(k - 1, 26); col = chr(65 + rem) + col
                address = col + str(ri + 1)
                cell = cells.get(address)
                if cell is None:
                    attributes = {"r": address}
                    if cells:
                        attributes["s"] = next(reversed(cells.values())).get("s", "0")
                    cell = ET.SubElement(row, tag("c"), attributes)
                if check:
                    stored = cell.find(tag("v"))
                    if cell.get("t") == "s" and stored is not None:
                        actual = shared_strings[int(stored.text)]
                    else:
                        actual = stored.text if stored is not None else "".join(cell.itertext())
                    if value and ri and rows[0][ci] in numeric:
                        assert float(actual) == float(value), (sheet_name, address)
                    else:
                        assert (actual or "") == value, (sheet_name, address)
                else:
                    for child in list(cell): cell.remove(child)
                    if value and ri and rows[0][ci] in numeric:
                        cell.attrib.pop("t", None)
                        ET.SubElement(cell, tag("v")).text = value
                    else:
                        cell.set("t", "inlineStr")
                        ET.SubElement(ET.SubElement(cell, tag("is")), tag("t")).text = value
        # Trim cells beyond the presentation table width.
        if not check:
            width = len(rows[0])
            for row in xml_rows:
                for cell in list(row.findall(tag("c"))):
                    letters = re.match(r"[A-Z]+", cell.get("r")).group()
                    number = 0
                    for letter in letters:
                        number = number * 26 + ord(letter) - 64
                    if number > width:
                        row.remove(cell)
            dimension = root.find(tag("dimension"))
            if dimension is not None and sheet_name == "SD8_Garratt_interactions":
                dimension.set("ref", "A1:I13")
            filter_node = root.find(tag("autoFilter"))
            if filter_node is not None and sheet_name == "SD8_Garratt_interactions":
                filter_node.set("ref", "A1:I13")
        if sheet_name == "SD1_ITP_experiments":
            last_column, number = "", len(rows[0])
            while number:
                number, rem = divmod(number - 1, 26)
                last_column = chr(65 + rem) + last_column
            table_ref = f"A1:{last_column}{len(rows)}"
            if check:
                assert root.find(tag("dimension")).get("ref") == table_ref, "SD1 dimensions"
            else:
                root.find(tag("dimension")).set("ref", table_ref)
                auto_filter = root.find(tag("autoFilter"))
                if auto_filter is not None:
                    auto_filter.set("ref", table_ref)
                columns = root.find(tag("cols"))
                if columns is None:
                    columns = ET.Element(tag("cols"))
                    root.insert(list(root).index(root.find(tag("sheetData"))), columns)
                for column in list(columns):
                    columns.remove(column)
                for index in range(len(rows[0])):
                    width = max(10, min(255, max(len(row[index]) for row in rows) + 3))
                    ET.SubElement(columns, tag("col"), {"min": str(index + 1), "max": str(index + 1),
                                  "width": str(width), "customWidth": "1", "bestFit": "1"})
        if sheet_name == "SD8_Garratt_interactions":
            for row in xml_rows[1:]:
                for cell in row.findall(tag("c")):
                    if re.match(r"[A-Z]+", cell.get("r")).group() not in {"F", "G", "H"}:
                        continue
                    index = int(cell.get("s", "0"))
                    if check:
                        assert cell_styles[index].get("numFmtId") == "2", "SD8 HR/CI precision"
                    elif cell_styles[index].get("numFmtId") != "2":
                        style = deepcopy(cell_styles[index])
                        style.set("numFmtId", "2")
                        style.set("applyNumberFormat", "1")
                        cell.set("s", str(len(cell_styles)))
                        cell_styles.append(style)
        if sheet_name in {"SD1_ITP_experiments", "SD7_external_eligibility"}:
            for row in xml_rows:
                if check:
                    assert float(row.get("ht", "15")) == 15, f"{sheet_name} row height"
                else:
                    row.set("ht", "15")
                    row.set("customHeight", "1")
                for cell in row.findall(tag("c")):
                    index = int(cell.get("s", "0"))
                    alignment = cell_styles[index].find(tag("alignment"))
                    if check:
                        assert alignment is None or alignment.get("wrapText", "0") in {"0", "false"}, f"{sheet_name} text wrapping"
                    elif alignment is None or alignment.get("wrapText") != "0":
                        if index not in no_wrap_styles:
                            style = deepcopy(cell_styles[index])
                            alignment = style.find(tag("alignment"))
                            if alignment is None:
                                alignment = ET.SubElement(style, tag("alignment"))
                            alignment.set("wrapText", "0")
                            style.set("applyAlignment", "1")
                            no_wrap_styles[index] = len(cell_styles)
                            cell_styles.append(style)
                        cell.set("s", str(no_wrap_styles[index]))
        if not check: parts[sheet] = xml_bytes(root)
    if not check:
        if no_wrap_styles:
            cell_styles.set("count", str(len(cell_styles)))
            parts["xl/styles.xml"] = xml_bytes(styles)
        with ZipFile(path, "w", ZIP_DEFLATED) as z:
            for name, raw in parts.items(): z.writestr(name, raw)


def build(check: bool = False) -> None:
    assessment = pd.read_csv(ROOT / "data/inputs/external_dataset_assessment.csv", keep_default_na=False)
    assert not assessment.duplicated(["study", "pmid"]).any()
    assert (assessment.source_url.eq("") | assessment.source_url.str.startswith("https://")).all()
    inventory = experiment_inventory()
    products = {
        "SD7_external_eligibility": csv_text(assessment),
        "SD1_ITP_experiments": csv_text(inventory),
    }
    interactions = pd.read_csv(OUT / "olfactory_cohort_interactions.csv")
    assert len(interactions) == 12
    assert not interactions.duplicated(["weight_age", "sex", "exposure"]).any()
    shown = interactions.copy()
    shown["weight_age"] = shown.weight_age.map({"12weeks": "12 weeks", "6months": "6 months", "12months": "12 months"})
    shown["sex"] = shown.sex.map({"F": "Female", "M": "Male"})
    shown["exposure"] = shown.exposure.map({"FU": "Adult-female odour", "MU": "Adult-male odour"})
    products["SD8_Garratt_interactions"] = csv_text(shown)
    sync_excel_tables(products, check)
    print(f"External datasets: {len(assessment)}; ITP source comparisons: {len(inventory)}; primary comparisons: {inventory.model.eq('Yes').sum()}")


def update_supplement_cells():
    ROOT = Path(__file__).absolute().parents[2]

    PATH = ROOT / "supplementary/Supplementary_Data_1-10.xlsx"

    TABLES = {"SD6_experiment_contrasts": "experiment_weight_changes.csv",
              "SD4_stratum_correlations": "weight_lifespan_by_experiment.csv"}

    original = PATH.read_bytes()
    workbook = load_workbook(BytesIO(original))
    changed_sheets = []
    count = 0
    for name, filename in TABLES.items():
        sheet = workbook[name]
        data = pd.read_csv(ROOT / "data/derived" / filename)
        if [c.value for c in sheet[1]] != data.columns.tolist():
            raise ValueError(f"Unexpected supplementary table structure: {name}")
        changed = sheet.max_row != len(data) + 1
        if sheet.max_row > len(data) + 1:
            sheet.delete_rows(len(data) + 2, sheet.max_row - len(data) - 1)
        for i, row in enumerate(data.itertuples(index=False, name=None), 2):
            for j, value in enumerate(row, 1):
                old = sheet.cell(i, j).value
                if isinstance(old, (float, int)) and isinstance(value, (float, int)):
                    equal = math.isclose(old, value, rel_tol=1e-12, abs_tol=1e-12)
                else:
                    equal = old == value
                if not equal:
                    sheet.cell(i, j).value = value
                    changed = True
                    count += 1
        if changed:
            changed_sheets.append(name)
    sheet = workbook["SD3_matched_quartiles"]
    sheet.cell(10, 2).value = "experiments"
    q = pd.read_csv(ROOT / "data/derived/matched_quartile_standardized_survival.csv")
    columns = ["quartile", "comparison_rows", "unique_mice", "mean_control_weight",
               "standardized_control_median", "standardized_treated_median", "gain_days",
               "gain_pct", "treatment_HR", "treatment_lo", "treatment_hi", "treatment_p"]
    for i, row in enumerate(q[columns].itertuples(index=False, name=None), 5):
        for j, value in enumerate(row, 1):
            sheet.cell(i, j).value = value
    trend = pd.read_csv(ROOT / "data/derived/matched_quartile_trend.csv").iloc[0]
    for j, key in enumerate(["experiments", "comparison_rows", "unique_mice", "HR_per_quartile", "lo", "hi", "p"], 2):
        sheet.cell(11, j).value = float(trend[key])
    sensitivity = pd.read_csv(ROOT / "data/derived/descriptive_sensitivities.csv")
    conventional = sensitivity[sensitivity.analysis.eq("weight-lowering conventional quartiles")].sort_values("bin")
    for j, value in enumerate(conventional.gain_pct, 2):
        sheet.cell(13, j).value = float(value)
    fifths = sensitivity[sensitivity.analysis.eq("weight-lowering standardized fifths")].sort_values("bin")
    for i, (label, column) in enumerate([("standardized_fifths_gain_pct", "gain_pct"),
                                        ("standardized_fifths_control_median_days", "control_median"),
                                        ("standardized_fifths_treated_median_days", "treated_median")], 14):
        sheet.cell(i, 1).value = label
        for j, value in enumerate(fifths[column], 2):
            sheet.cell(i, j).value = float(value)
    changed_sheets.append("SD3_matched_quartiles")
    readme = workbook["README"]
    contrasts = len(pd.read_csv(ROOT / "data/derived/experiment_weight_changes.csv"))
    strata = len(pd.read_csv(ROOT / "data/derived/weight_lifespan_by_experiment.csv"))
    for row in readme.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = re.sub(r"\d+ cohort x sex x experiment", f"{contrasts} cohort x sex x experiment", cell.value)
                cell.value = re.sub(r"\d+ cohort x experiment x sex strata", f"{strata} cohort x experiment x sex strata", cell.value)
                cell.value = re.sub(r"\d+ male weight-lowering experiments \([\d,]+ unique mice\)",
                                    f"{int(trend.experiments)} male weight-lowering experiments ({int(trend.unique_mice):,} unique mice)", cell.value)
    description = readme.cell(3, 3)
    exclusion_note = " Protandim is retained in broader intervention analyses but excluded from single-compound analyses because it is a multicomponent formulation."
    if exclusion_note.strip() not in description.value:
        description.value += exclusion_note
    control_records = pd.read_excel(PATH, sheet_name="SD10_control_lifespan")
    itp = control_records.itp_status.astype(str).str.upper().eq("ITP")
    count_by_group = lambda sex, inside: int((control_records.sex.eq(sex) & (itp if inside else ~itp)).sum())
    description = (f"Control-lifespan corpus: {len(control_records)} experiments from {control_records.pmid.nunique()} studies "
                   f"({count_by_group('Female', True)} ITP and {count_by_group('Female', False)} non-ITP female experiments; "
                   f"{count_by_group('Male', True)} ITP and {count_by_group('Male', False)} non-ITP male experiments; "
                   f"{int(control_records.sex.eq('Pooled').sum())} pooled-sex experiments). "
                   "Used for Extended Data Table 7 and Fig. 3b,c. UM-HET3 studies outside ITP are classified as non-ITP.")
    for row in readme.iter_rows():
        if len(row) >= 3 and row[1].value == "SD10_control_lifespan":
            row[2].value = description
    changed_sheets.append("README")
    if changed_sheets:
        rendered = BytesIO()
        workbook.save(rendered)
        result = BytesIO()
        # Copy only the updated sheet parts; retain all other original ZIP parts.
        with zipfile.ZipFile(BytesIO(original)) as source, zipfile.ZipFile(rendered) as edited:
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            rels = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
            targets = {r.attrib["Id"]: r.attrib["Target"] for r in rels}
            book = ET.fromstring(source.read("xl/workbook.xml"))
            edited_rels = ET.fromstring(edited.read("xl/_rels/workbook.xml.rels"))
            edited_targets = {r.attrib["Id"]: r.attrib["Target"] for r in edited_rels}
            edited_book = ET.fromstring(edited.read("xl/workbook.xml"))
            edited_parts = {}
            for sheet in edited_book.findall("s:sheets/s:sheet", ns):
                rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                target = edited_targets[rid]
                edited_parts[sheet.attrib["name"]] = target.lstrip("/") if target.startswith("/") else "xl/" + target
            parts = {}
            for sheet in book.findall("s:sheets/s:sheet", ns):
                if sheet.attrib["name"] in changed_sheets:
                    rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                    target = targets[rid]
                    source_part = target.lstrip("/") if target.startswith("/") else "xl/" + target
                    parts[source_part] = edited_parts[sheet.attrib["name"]]
            with zipfile.ZipFile(result, "w") as destination:
                for info in source.infolist():
                    destination.writestr(info, edited.read(parts[info.filename]) if info.filename in parts
                                         else source.read(info.filename))
        PATH.write_bytes(result.getvalue())
    workbook.close()
    print(f"Supplementary results: {count} cells updated.")


RESULT_TABLES = {'adult_weight_qc.csv': 'adult_weight_qc.csv',
 'experiment_interactions.csv': 'experiment_interactions.csv',
 'comparator_gradients.csv': 'comparator_gradients.csv',
 'comparator_strata.csv': 'comparator_state.csv',
 'compound_survival.csv': 'compound_descriptive.csv',
 'control_lifespan_comparison.csv': 'control_lifespan_ceiling.csv',
 'control_weight_gradients.csv': 'control_weight_gradient_by_age.csv',
 'quartile_sensitivities.csv': 'descriptive_sensitivities.csv',
 'restriction_quartiles.csv': 'drido_pooled_cr_quartiles.csv',
 'adult_weight_change.csv': 'early_weight_change.csv',
 'external_survival_context.csv': 'external_sex_survival_context.csv',
 'maternal_adjustment.csv': 'maternal_adjustment.csv',
 'matched_quartile_survival.csv': 'matched_quartile_standardized_survival.csv',
 'olfactory_interactions.csv': 'olfactory_cohort_interactions.csv',
 'restricted_mean_survival.csv': 'quartile_rmst.csv',
 'restricted_mean_contrasts.csv': 'quartile_rmst_contrasts.csv',
 'sex_survival_by_cohort.csv': 'sex_survival_by_cohort.csv',
 'sex_survival_summary.csv': 'sex_survival_context.csv',
 'cross_study_weight_response.csv': 'supplementary_analyses.csv',
 'weight_response_quartiles.csv': 'weight_change_quartiles.csv'}

RESULT_GROUPS = {'treatment_moderation.csv': [['primary', 'primary_moderation.csv'],
                              ['weight_class', 'weight_lowering_split.csv'],
                              ['landmark', 'successful_landmarks.csv'],
                              ['landmark_lag90', 'successful_landmarks_lag90.csv'],
                              ['landmark_leave_cohort_out', 'successful_landmarks_loo12.csv']],
 'site_consistency.csv': [['site_specific', 'site_moderation.csv'],
                          ['heterogeneity', 'site_heterogeneity.csv']],
 'strain_comparison.csv': [['founder_weights', 'founder_strain_weight_comparison.csv'],
                           ['strain_context', 'strain_context_summary.csv']],
 'developmental_weights.csv': [['puberty_moderation', 'puberty_moderation.csv'],
                               ['joint_puberty_adult', 'puberty_joint.csv'],
                               ['developmental_gradients', 'developmental_figure_estimates.csv']],
 'restriction_models.csv': [['treatment_by_quarter', 'drido_pooled_cr_interaction.csv'],
                            ['weight_mortality', 'drido_weight_mortality_by_group.csv']],
 'survival_curves.csv': [['all_weight_lowering', 'matched_quartile_survival_curves.csv'],
                         ['rapamycin', 'rapamycin_quartile_survival_curves.csv']],
 'weight_response_models.csv': [['continuous', 'weight_change_slopes.csv'],
                                ['quartile_trend', 'weight_change_trend.csv'],
                                ['comparator_adjustment', 'weight_change_vs_comparator.csv']],
 'olfactory_weight_models.csv': [['model_based', 'olfactory_cohort_weight_mortality.csv']],
 'quartile_tests.csv': [['matched_quartiles', 'matched_quartile_trend.csv'],
                        ['common_control_thresholds', 'common_threshold_quartile_trend.csv'],
                        ['compound_common_interaction', 'compound_common_interaction.csv']]}


def verify_derived_supplement_sheets():
    """Compare the shipped presentation tables with the regenerated calculations."""
    book = load_workbook(SUP / "Supplementary_Data_1-10.xlsx", read_only=True, data_only=True)
    try:
        def equal(actual, expected):
            if pd.isna(expected):
                return actual is None or pd.isna(actual)
            if isinstance(expected, (float, int)) and isinstance(actual, (float, int)):
                return math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-12)
            return actual == expected
        for sheet_name, filename in [("SD6_experiment_contrasts", "experiment_weight_changes.csv"),
                                      ("SD4_stratum_correlations", "weight_lifespan_by_experiment.csv")]:
            frame = pd.read_csv(OUT / filename)
            sheet = book[sheet_name]
            assert [c.value for c in sheet[1]] == frame.columns.tolist(), sheet_name
            assert sheet.max_row == len(frame) + 1, sheet_name
            for i, values in enumerate(frame.itertuples(index=False, name=None), 2):
                for j, value in enumerate(values, 1):
                    assert equal(sheet.cell(i, j).value, value), (sheet_name, i, j)
        frame = pd.read_csv(OUT / "matched_quartile_standardized_survival.csv")
        columns = ["quartile", "comparison_rows", "unique_mice", "mean_control_weight",
                   "standardized_control_median", "standardized_treated_median", "gain_days",
                   "gain_pct", "treatment_HR", "treatment_lo", "treatment_hi", "treatment_p"]
        sheet = book["SD3_matched_quartiles"]
        for i, values in enumerate(frame[columns].itertuples(index=False, name=None), 5):
            for j, value in enumerate(values, 1):
                assert equal(sheet.cell(i, j).value, value), (sheet.title, i, j)
        trend = pd.read_csv(OUT / "matched_quartile_trend.csv").iloc[0]
        for j, key in enumerate(["experiments", "comparison_rows", "unique_mice", "HR_per_quartile", "lo", "hi", "p"], 2):
            assert equal(sheet.cell(11, j).value, trend[key]), (sheet.title, 11, j)
        sensitivity = pd.read_csv(OUT / "descriptive_sensitivities.csv")
        conventional = sensitivity[sensitivity.analysis.eq("weight-lowering conventional quartiles")].sort_values("bin")
        for j, value in enumerate(conventional.gain_pct, 2):
            assert equal(sheet.cell(13, j).value, value), (sheet.title, 13, j)
        fifths = sensitivity[sensitivity.analysis.eq("weight-lowering standardized fifths")].sort_values("bin")
        for i, (label, column) in enumerate([("standardized_fifths_gain_pct", "gain_pct"),
                                            ("standardized_fifths_control_median_days", "control_median"),
                                            ("standardized_fifths_treated_median_days", "treated_median")], 14):
            assert sheet.cell(i, 1).value == label, (sheet.title, i, 1)
            for j, value in enumerate(fifths[column], 2):
                assert equal(sheet.cell(i, j).value, value), (sheet.title, i, j)
    finally:
        book.close()


def compare_result_table(released, generated, name):
    """Compare labels and counts exactly, allowing relative floating-point error."""
    actual = pd.read_csv(io.BytesIO(released), dtype=str, keep_default_na=False)
    expected = pd.read_csv(io.BytesIO(generated), dtype=str, keep_default_na=False)
    assert actual.columns.tolist() == expected.columns.tolist() and actual.shape == expected.shape, name
    exact_name = re.compile(r"(?:^|_)(?:id|n|mice|deaths|events|experiments|cohorts|clusters|replicates|rows|cells|sites|compounds|df)(?:$|_)")
    for column in expected:
        a, b = actual[column], expected[column]
        if a.equals(b):
            continue
        missing = ("", "NA")
        assert a.isin(missing).equals(b.isin(missing)), (name, column, "missing values")
        present = ~b.isin(missing)
        if a[present].equals(b[present]):
            continue
        if exact_name.search(column):
            assert a[present].equals(b[present]), (name, column)
        else:
            try:
                x = pd.to_numeric(a[present], errors="raise").to_numpy()
                y = pd.to_numeric(b[present], errors="raise").to_numpy()
            except (ValueError, TypeError):
                raise AssertionError((name, column, "labels differ")) from None
            assert np.allclose(x, y, rtol=1e-9, atol=0, equal_nan=False), (name, column)


def export_result_tables(check=False):
    """Write or validate compact CSV result tables."""
    derived = ROOT / "data/derived"
    results = ROOT / "data/outputs"
    if not check:
        results.mkdir(parents=True, exist_ok=True)
    products = {}
    for name, source in RESULT_TABLES.items():
        if name == "olfactory_interactions.csv":
            frame = pd.read_csv(derived / source)
            frame["weight_age"] = frame.weight_age.map({"12weeks": "12 weeks", "6months": "6 months", "12months": "12 months"})
            frame["sex"] = frame.sex.map({"F": "Female", "M": "Male"})
            frame["exposure"] = frame.exposure.map({"FU": "Adult-female odour", "MU": "Adult-male odour"})
            products[name] = frame.to_csv(index=False, lineterminator="\n").encode("utf-8")
        else:
            products[name] = (derived / source).read_bytes().replace(b"\r\n", b"\n")
    for name, sources in RESULT_GROUPS.items():
        columns, rows = ["model_set"], []
        for label, source in sources:
            with (derived / source).open(encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                columns.extend(column for column in reader.fieldnames if column not in columns)
                rows.extend({"model_set": label, **row} for row in reader)
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=columns, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
        products[name] = buffer.getvalue().encode("utf-8")
    for name, raw in products.items():
        path = results / name
        if check:
            compare_result_table(path.read_bytes(), raw, name)
        else:
            path.write_bytes(raw)
    assert {p.name for p in results.glob("*.csv")} == set(products)
    if not check:
        print(f"Exported {len(products)} manuscript result tables.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true", help="Check result files without overwriting them")
    parser.add_argument("--rebuild-supplement", action="store_true", help="Update calculated workbook sheets")
    args = parser.parse_args()
    if args.rebuild_supplement:
        update_supplement_cells()
        build(check=False)
    else:
        verify_derived_supplement_sheets()
        build(check=True)
    export_result_tables(check=args.check)
